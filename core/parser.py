"""Resume Upload & Parsing Module.

Validates uploads and extracts raw text from PDF, DOCX, TXT and image files.
Scanned PDFs (pages with no text layer) are rasterised and sent through the OCR module.
"""
import io
import re
import time
import zipfile
from pathlib import Path

import config
from core import nlp, ocr

MAGIC = {
    ".pdf": [b"%PDF"],
    ".docx": [b"PK\x03\x04"],
    ".png": [b"\x89PNG"],
    ".jpg": [b"\xff\xd8\xff"],
    ".jpeg": [b"\xff\xd8\xff"],
    ".bmp": [b"BM"],
    ".tif": [b"II*\x00", b"MM\x00*"],
    ".tiff": [b"II*\x00", b"MM\x00*"],
}


class InvalidFileError(ValueError):
    pass


def safe_name(name: str) -> str:
    return re.sub(r"[^\w.\- ]", "_", Path(name).name)[:120]


def validate_file(name: str, data: bytes):
    """File type/size validation: extension whitelist, size cap and magic-byte check."""
    ext = Path(name).suffix.lower()
    if ext not in config.ALLOWED_EXT:
        raise InvalidFileError(f"'{name}': unsupported file type {ext or '(none)'}.")
    if not data:
        raise InvalidFileError(f"'{name}' is empty.")
    if len(data) > config.MAX_FILE_MB * 1024 * 1024:
        raise InvalidFileError(f"'{name}' exceeds the {config.MAX_FILE_MB} MB limit.")
    if ext in MAGIC and not any(data.startswith(m) for m in MAGIC[ext]):
        raise InvalidFileError(f"'{name}': file content does not match its {ext} extension.")
    if ext == ".txt":
        try:
            data.decode("utf-8")
        except UnicodeDecodeError:
            try:
                data.decode("latin-1")
            except UnicodeDecodeError as exc:
                raise InvalidFileError(f"'{name}' is not a readable text file.") from exc


def _pdf(data: bytes) -> dict:
    import fitz  # PyMuPDF

    texts, scanned_pages, ocr_results = [], 0, []
    with fitz.open(stream=data, filetype="pdf") as doc:
        if doc.page_count == 0:
            raise InvalidFileError("PDF has no pages.")
        for page in doc:
            txt = page.get_text("text")
            if len(txt.strip()) >= config.SCANNED_PDF_MIN_CHARS:
                texts.append(txt)
                continue
            scanned_pages += 1
            pix = page.get_pixmap(dpi=300)
            img = ocr.decode_image(pix.tobytes("png"))
            res = ocr.ocr_image(img)
            ocr_results.append(res)
            texts.append(res["text"])
    if scanned_pages:
        return {
            "text": "\n".join(texts),
            "resume_type": "scanned",
            "ocr_confidence": round(sum(r["ocr_confidence"] for r in ocr_results) / len(ocr_results), 1),
            "cv_layout_score": round(sum(r["layout_score"] for r in ocr_results) / len(ocr_results), 1),
        }
    return {"text": "\n".join(texts), "resume_type": "text"}


def _docx(data: bytes) -> dict:
    try:
        import docx

        document = docx.Document(io.BytesIO(data))
        parts = [p.text for p in document.paragraphs]
        for table in document.tables:
            for row in table.rows:
                parts.append(" | ".join(cell.text for cell in row.cells))
        text = "\n".join(parts)
    except ImportError:
        with zipfile.ZipFile(io.BytesIO(data)) as z:
            xml = z.read("word/document.xml").decode("utf-8", errors="ignore")
        xml = re.sub(r"</w:p>", "\n", xml)
        text = re.sub(r"<[^>]+>", "", xml)
    return {"text": text, "resume_type": "text"}


def _image(data: bytes) -> dict:
    res = ocr.ocr_image(ocr.decode_image(data))
    return {
        "text": res["text"],
        "resume_type": "scanned",
        "ocr_confidence": res["ocr_confidence"],
        "cv_layout_score": res["layout_score"],
    }


def extract_text(name: str, data: bytes) -> dict:
    validate_file(name, data)
    ext = Path(name).suffix.lower()
    if ext == ".pdf":
        return _pdf(data)
    if ext == ".docx":
        return _docx(data)
    if ext == ".txt":
        try:
            return {"text": data.decode("utf-8"), "resume_type": "text"}
        except UnicodeDecodeError:
            return {"text": data.decode("latin-1"), "resume_type": "text"}
    return _image(data)


def process_resume(name: str, data: bytes) -> dict:
    """Full pipeline: validate -> extract (OCR if needed) -> NLP analysis. Returns a DB-ready dict."""
    start = time.perf_counter()
    extracted = extract_text(name, data)
    if len(extracted["text"].strip()) < 20:
        raise InvalidFileError(f"'{name}': no readable text could be extracted.")
    info = nlp.analyze_resume(extracted["text"])
    content_score = info.pop("structure_score")
    if extracted["resume_type"] == "scanned":
        layout = round(0.5 * content_score + 0.5 * extracted["cv_layout_score"], 1)
    else:
        layout = content_score
    info.update(
        file_name=safe_name(name),
        resume_type=extracted["resume_type"],
        ocr_confidence=extracted.get("ocr_confidence"),
        layout_score=layout,
        processing_time=round(time.perf_counter() - start, 3),
    )
    if not info["name"]:
        info["name"] = Path(name).stem.replace("_", " ").replace("-", " ").title()
    return info
